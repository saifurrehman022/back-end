import os
import re
import logging
from typing import List, Tuple, Optional
import faiss
from sentence_transformers import SentenceTransformer
from PyPDF2 import PdfReader
from docx import Document
import pytesseract
from PIL import Image
import io
import openpyxl
import pandas as pd
from duckduckgo_search import DDGS
from fastapi import UploadFile

logger = logging.getLogger(__name__)

_EMBED_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
_embedder: Optional[SentenceTransformer] = None

def _get_embedder() -> SentenceTransformer:
    global _embedder
    if _embedder is None:
        logger.info(f"Loading embedding model: {_EMBED_MODEL_NAME}")
        _embedder = SentenceTransformer(_EMBED_MODEL_NAME)
    return _embedder

# Enhanced File Extraction
def extract_text(file: UploadFile) -> str:
    ext = os.path.splitext(file.filename)[1].lower()
    content = file.file.read()
    file_bytes = io.BytesIO(content)
    if ext == ".pdf":
        try:
            reader = PdfReader(file_bytes)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            logger.error(f"PDF extract failed: {e}")
            return ""
    elif ext == ".docx":
        try:
            doc = Document(file_bytes)
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception as e:
            logger.error(f"DOCX extract failed: {e}")
            return ""
    elif ext in [".xlsx", ".xls"]:
        try:
            wb = openpyxl.load_workbook(file_bytes, read_only=True, data_only=True)
            text = []
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    text.append(" ".join(str(cell) for cell in row if cell is not None))
            return "\n".join(text)
        except Exception as e:
            logger.error(f"Excel extract failed: {e}")
            return ""
    elif ext == ".csv":
        try:
            df = pd.read_csv(file_bytes)
            return df.to_string()
        except Exception as e:
            logger.error(f"CSV extract failed: {e}")
            return ""
    elif ext in [".jpg", ".jpeg", ".png", ".gif"]:  # OCR for images
        try:
            img = Image.open(file_bytes)
            return pytesseract.image_to_string(img)
        except Exception as e:
            logger.error(f"Image OCR failed: {e}")
            return ""
    else:  # Fallback text
        try:
            return content.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"Text extract failed: {e}")
            return ""

def clean_text(text: str) -> str:
    t = re.sub(r"[ \t]+", " ", text)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()

def chunk_text(text: str, max_tokens: int = 400, overlap: int = 50) -> List[str]:
    text = clean_text(text)
    if not text:
        return []
    words = text.split()
    chunks, start = [], 0
    while start < len(words):
        end = min(len(words), start + max_tokens)
        chunk = " ".join(words[start:end]).strip()
        if chunk:
            chunks.append(chunk)
        if end == len(words):
            break
        start = max(0, end - overlap)
    return chunks

class RagIndex:
    def __init__(self, index: faiss.IndexFlatIP, dim: int, chunks: List[str]):
        self.index = index
        self.dim = dim
        self.chunks = chunks

def build_faiss_index(chunks: List[str]) -> RagIndex:
    emb = _get_embedder()
    vectors = emb.encode(chunks, convert_to_numpy=True, normalize_embeddings=True)
    dim = vectors.shape[1]
    index = faiss.IndexFlatIP(dim)
    index.add(vectors)
    return RagIndex(index=index, dim=dim, chunks=chunks)

def search(index: RagIndex, query: str, top_k: int = 6) -> List[Tuple[str, float]]:
    emb = _get_embedder()
    q = emb.encode([query], convert_to_numpy=True, normalize_embeddings=True)
    D, I = index.index.search(q, top_k)
    hits = []
    for score, idx in zip(D[0], I[0]):
        if idx == -1:
            continue
        hits.append((index.chunks[idx], float(score)))
    return hits

def build_context_from_files(files: List[UploadFile], prompt: str, top_k: int = 6) -> str:
    all_text = []
    for file in files:
        txt = extract_text(file)
        if txt:
            all_text.append(txt)
        file.file.seek(0)  # Reset
    big_text = "\n\n".join(all_text)
    chunks = chunk_text(big_text, max_tokens=450, overlap=80)
    if not chunks:
        return ""
    idx = build_faiss_index(chunks)
    hits = search(idx, prompt, top_k=top_k)
    context_sections = [f"[DOC#{i} score={score:.3f}]\n{chunk}" for i, (chunk, score) in enumerate(hits, 1)]
    return "\n\n".join(context_sections)

# Web search tool
def web_search(query: str) -> str:
    try:
        with DDGS() as ddgs:
            results = [r for r in ddgs.text(query, max_results=5)]
        sections = [f"[WEB#{i}] Title: {r['title']}\nSnippet: {r['body']}\nURL: {r['href']}" for i, r in enumerate(results, 1)]
        return "\n\n".join(sections) if sections else "No results found."
    except Exception as e:
        logger.error(f"Web search failed: {e}")
        return "Web search error."